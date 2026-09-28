"""Language-pair auto-detection for bilingual imports (issue #200).

The DOCX files are built here with python-docx the way Trados and CafeTran
write them: a table whose source and target runs carry Word language tags.
"""

from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from modules.bilingual_lang_detect import docx_column_languages, header_language, known_code
from modules.cafetran_docx_handler import CafeTranDOCXHandler
from modules.dejavurtf_handler import DejaVuRTFHandler
from modules.trados_docx_handler import TradosDOCXHandler

TRADOS_HEADER = ["Segment ID", "Segment status", "Source segment", "Target segment"]


def _tagged_run(cell, text, lang=None, east_asia=None):
    run = cell.paragraphs[0].add_run(text)
    if lang or east_asia:
        rpr = run._r.get_or_add_rPr()
        el = OxmlElement('w:lang')
        if lang:
            el.set(qn('w:val'), lang)
        if east_asia:
            el.set(qn('w:eastAsia'), east_asia)
        rpr.append(el)


def _make_docx(path, header, rows):
    """rows: [(cells_text_and_langs)], each cell a (text, lang, east_asia) tuple or a str."""
    doc = Document()
    table = doc.add_table(rows=1, cols=len(header))
    for i, h in enumerate(header):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, cell_spec in enumerate(row):
            if isinstance(cell_spec, str):
                cells[i].text = cell_spec
            else:
                _tagged_run(cells[i], *cell_spec)
    doc.save(path)
    return str(path)


def _trados(path, rows):
    handler = TradosDOCXHandler()
    assert handler.load(_make_docx(path, TRADOS_HEADER, rows))
    return handler.detect_language_pair()


def test_trados_review_docx_pair_comes_from_each_columns_tags(tmp_path):
    rows = [("1", "Translated", ("Die Pumpe", "de-DE"), ("The pump", "en-GB")),
            ("2", "Draft", ("Das Gehäuse", "de-DE"), ("", "en-GB")),
            ("3", "Draft", ("Ventil", "de-DE"), ("", None))]
    assert _trados(tmp_path / "t.docx", rows) == ("de-DE", "en-GB")


def test_same_language_in_both_columns_is_not_trusted(tmp_path):
    rows = [("1", "Draft", ("Pumpe", "en-US"), ("Pump", "en-US"))]
    assert _trados(tmp_path / "t.docx", rows) == (None, None)


def test_an_untagged_target_column_gives_a_partial_result(tmp_path):
    rows = [("1", "Draft", ("Pumpe", "de-DE"), "")]
    assert _trados(tmp_path / "t.docx", rows) == ("de-DE", None)


def test_untagged_file_detects_nothing(tmp_path):
    rows = [("1", "Draft", "Pumpe", "Pump")]
    assert _trados(tmp_path / "t.docx", rows) == (None, None)


def test_cjk_text_uses_the_east_asian_language_slot(tmp_path):
    rows = [("1", "Draft", ("The pump", "en-US"), ("ポンプ", "en-US", "ja-JP"))]
    assert _trados(tmp_path / "t.docx", rows) == ("en-US", "ja-JP")


def test_placeholder_language_values_are_ignored(tmp_path):
    rows = [("1", "Draft", ("x", "x-none"), ("y", "nl-NL"))]
    assert _trados(tmp_path / "t.docx", rows) == (None, "nl-NL")


def test_cafetran_header_codes_win_then_run_tags(tmp_path):
    handler = CafeTranDOCXHandler()
    assert handler.load(_make_docx(tmp_path / "c.docx", ["ID", "EN-GB", "NL-NL", "Notes", "*"],
                                   [("1", "Pump", "", "", "")]))
    assert handler.detect_language_pair() == ("en-GB", "nl-NL")

    handler = CafeTranDOCXHandler()
    assert handler.load(_make_docx(tmp_path / "c2.docx", ["ID", "manual.docx", "manual.docx", "Notes", "*"],
                                   [("1", ("Pumpe", "de-DE"), ("Pomp", "nl-BE"), "", "")]))
    assert handler.detect_language_pair() == ("de-DE", "nl-BE")


def test_codes_and_headers_are_not_read_out_of_ordinary_words():
    assert known_code("de-de") == "de-DE"
    assert known_code("nld") == "nl"
    assert known_code("Notes") is None and known_code("x-none") is None and known_code("") is None
    assert header_language("Dutch (Belgium)").startswith("nl")
    assert header_language("French") == "fr"
    assert header_language("Source segment") is None
    assert header_language("manual.docx") is None


def test_column_detection_survives_a_short_table():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    assert docx_column_languages(table, 2, 3) == (None, None)


def test_dejavu_pair_is_only_reported_when_detected():
    handler = DejaVuRTFHandler()
    assert handler.detect_language_pair() == (None, None)  # placeholders, not a detection
    handler.languages_detected = True
    handler.source_lang, handler.target_lang = "German (CH)", "English (UK)"
    assert handler.detect_language_pair() == ("de", "en")
    handler.target_lang = "Unknown (9999)"
    assert handler.detect_language_pair() == ("de", None)


@pytest.mark.parametrize("name, code", [("German", 1031), ("Portuguese", 1046), ("pt", 1046),
                                        ("Spanish (MX)", 2058), ("Klingon", None)])
def test_dejavu_export_language_code_falls_back_to_the_base_language(name, code):
    assert DejaVuRTFHandler()._get_rtf_lang_code(name) == code
