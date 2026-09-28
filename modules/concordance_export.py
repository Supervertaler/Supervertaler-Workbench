"""
Concordance Export
==================

Writes concordance (SuperLookup TM) search results to Excel or CSV, so hits can
be kept, shared or used to document inconsistencies in a TM (issue #170).

Kept free of Qt so it can be tested on its own; the SuperLookup tab collects
the rows (in the order the table currently shows them) and picks the path.
"""

import csv
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional


def export_concordance_results(rows: List[Dict], path: str, query: str = "",
                               source_label: str = "Source",
                               target_label: str = "Target") -> int:
    """Write concordance hits to ``path``; the format follows its extension.

    Args:
        rows: dicts with 'source', 'target' and 'tm_name' keys.
        path: destination; ``.xlsx`` writes an Excel workbook, anything else CSV.
        query: the search term, recorded in the workbook's "Search" sheet.
        source_label / target_label: column headings, e.g. the language names.

    Returns:
        The number of hits written.
    """
    header = [source_label or "Source", target_label or "Target", "TM"]
    data = [[r.get('source', '') or '', r.get('target', '') or '', r.get('tm_name', '') or '']
            for r in rows]

    if Path(path).suffix.lower() == '.xlsx':
        _write_xlsx(path, header, data, query)
    else:
        # utf-8-sig so Excel detects the encoding and accented letters survive.
        with open(path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(header)
            writer.writerows(data)
    return len(data)


def _write_xlsx(path: str, header: List[str], data: List[List[str]], query: Optional[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Concordance"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    wrap = Alignment(wrap_text=True, vertical='top')
    for values in data:
        ws.append(values)
        for cell in ws[ws.max_row]:
            # openpyxl stores any string beginning with "=" as a formula, which
            # would turn TM text into a broken (or live) formula in Excel.
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.data_type = 's'
            cell.alignment = wrap

    ws.column_dimensions['A'].width = 60
    ws.column_dimensions['B'].width = 60
    ws.column_dimensions['C'].width = 25
    ws.freeze_panes = 'A2'

    info = wb.create_sheet("Search")
    for label, value in (("Search term", query or ""),
                         ("Hits", len(data)),
                         ("Exported", datetime.now().strftime("%Y-%m-%d %H:%M"))):
        info.append([label, value])
        info.cell(row=info.max_row, column=1).font = Font(bold=True)
        value_cell = info.cell(row=info.max_row, column=2)
        if isinstance(value_cell.value, str) and value_cell.value.startswith('='):
            value_cell.data_type = 's'
    info.column_dimensions['A'].width = 14
    info.column_dimensions['B'].width = 50

    wb.save(path)
